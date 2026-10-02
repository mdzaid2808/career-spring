"""CareerSpring — hackathon career counselor prototype.
Run: streamlit run app.py
"""
import os
from pathlib import Path

import streamlit as st
from dotenv import load_dotenv
from groq import APIError, Groq

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None
try:
    from docx import Document
except ImportError:
    Document = None

load_dotenv(Path(__file__).with_name(".env"), override=False)
st.set_page_config(page_title="CareerSpring | Career Studio", page_icon="✦", layout="wide")


def secret(name: str, default: str = "") -> str:
    value = os.getenv(name, "").strip()
    if value:
        return value
    try:
        return str(st.secrets.get(name, default)).strip()
    except Exception:
        return default


API_KEY = secret("GROQ_API_KEY")
MODEL = secret("GROQ_MODEL", "openai/gpt-oss-20b")

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600;700&family=Manrope:wght@400;500;600;700;800&display=swap');
:root { color-scheme: dark; }
.stApp { background: radial-gradient(ellipse at 78% 0%, #13342e 0%, #0b121a 38%, #080d13 75%); color:#e9f2f2; }
[data-testid="stHeader"] {background:transparent;}
[data-testid="stSidebar"] {background:#0d151d;border-right:1px solid #24333b;}
.block-container {max-width:1320px;padding:2.2rem 2.4rem 4rem;}
h1,h2,h3 {font-family:'Manrope',sans-serif;letter-spacing:-.7px;color:#eef8f5!important;}
p,li, label {font-family:'DM Sans',sans-serif;line-height:1.65;}
.hero {padding:30px 34px;border-radius:24px;border:1px solid #284941;background:linear-gradient(120deg,#142b2a,#101b24 68%);margin-bottom:22px;}
.kicker {color:#84ddc0;text-transform:uppercase;letter-spacing:2px;font-size:11px;font-weight:700;}
.hero h1 {font-size:40px;margin:10px 0 6px;}
.hero p {color:#a8bec0;margin:0;max-width:760px;}
.pill {display:inline-block;border:1px solid #315148;border-radius:20px;padding:6px 10px;color:#a6e2ce;font-size:11px;margin:12px 5px 0 0;}
.card {background:#111c25;border:1px solid #253742;border-radius:18px;padding:20px 22px;margin-bottom:12px;}
.card-title {font-size:17px;font-weight:700;color:#e5f5ef;}
.card-sub {font-size:13px;color:#9cb0bb;margin-top:5px;}
div.stButton > button {border-radius:11px;border:1px solid #315147;background:#18332d;color:#dff7ee;font-weight:600;}
div.stButton > button:hover {border-color:#78d9ba;color:white;}
[data-testid="stFileUploader"] {border:1px dashed #426459;border-radius:14px;background:#101a21;}
[data-testid="stChatMessage"] {border:1px solid #22333b;border-radius:16px;background:#101a22;}
hr {border-color:#26363e;}
small,.muted {color:#9aadb4;}
</style>
""", unsafe_allow_html=True)

if "career_messages" not in st.session_state:
    st.session_state.career_messages = []
if "interview_messages" not in st.session_state:
    st.session_state.interview_messages = []
if "interview_question" not in st.session_state:
    st.session_state.interview_question = 0

with st.sidebar:
    st.markdown("## ✦ CareerSpring")
    st.caption("AI CAREER STUDIO · HACKATHON PROTOTYPE")
    st.divider()
    language = st.selectbox("Reply language", ["English", "Hinglish", "Hindi"], index=1)
    education = st.text_input("Your education / year", placeholder="e.g., B.Tech CSE · 2nd year")
    interests = st.text_input("Interests", placeholder="e.g., AI, design, healthcare")
    skills = st.text_input("Skills you already have", placeholder="e.g., Python, communication")
    st.info("Career guidance only. Suggestions are starting points—verify current salary, openings and course details independently.")
    st.caption("Resume text is processed for this session and sent to the AI provider when you request analysis. Avoid uploading sensitive identifiers.")
    if st.button("Clear chat history", use_container_width=True):
        st.session_state.career_messages = []
        st.session_state.interview_messages = []
        st.session_state.interview_question = 0
        st.rerun()

st.markdown("""
<div class="hero">
  <div class="kicker">Find your next step</div>
  <h1>Your career, with a clearer direction.</h1>
  <p>Explore roles, strengthen your resume, build a practical learning plan, and rehearse interviews in one focused workspace.</p>
  <span class="pill">Career guide</span><span class="pill">Resume review</span><span class="pill">Skill roadmap</span><span class="pill">Mock interview</span>
</div>
""", unsafe_allow_html=True)

if not API_KEY:
    st.warning("Groq key is not configured yet. Add GROQ_API_KEY in Streamlit Secrets to enable AI features.")

client = Groq(api_key=API_KEY, timeout=60, max_retries=0) if API_KEY else None
profile = f"Education: {education or 'not provided'}\nInterests: {interests or 'not provided'}\nCurrent skills: {skills or 'not provided'}"


def complete(system_prompt: str, user_content: str, history=None, max_tokens=1400) -> str:
    if not client:
        raise RuntimeError("Groq API key is missing. Add GROQ_API_KEY in app Secrets.")
    messages = [{"role": "system", "content": system_prompt}]
    if history:
        messages.extend(history[-10:])
    messages.append({"role": "user", "content": user_content})
    result = client.chat.completions.create(
        model=MODEL,
        messages=messages,
        temperature=0.35,
        max_completion_tokens=max_tokens,
    )
    answer = result.choices[0].message.content
    if not answer or not answer.strip():
        raise RuntimeError("The AI returned an empty response. Try again.")
    return answer.strip()


def show_error(error):
    code = getattr(error, "status_code", None)
    if code == 401:
        st.error("Groq key was rejected. Check your app Secrets.")
    elif code == 429:
        st.error("Groq rate limit reached. Wait briefly and try again.")
    elif code in (400, 403, 404):
        st.error("The selected model may be unavailable for this account. Check GROQ_MODEL in Secrets.")
    else:
        st.error("Could not reach the AI service. Check the deployment logs and try again.")


tab_chat, tab_resume, tab_roadmap, tab_interview = st.tabs([
    "💬 Career Counselor", "📄 Resume Analyzer", "🧭 Learning Roadmap", "🎤 Mock Interview"
])

with tab_chat:
    st.subheader("Talk through your career options")
    st.caption("Tell the counselor what you enjoy, what you are good at, and what constraints matter to you.")
    examples = st.columns(3)
    with examples[0]:
        st.markdown('<div class="card"><div class="card-title">Explore roles</div><div class="card-sub">Compare paths that fit your strengths.</div></div>', unsafe_allow_html=True)
    with examples[1]:
        st.markdown('<div class="card"><div class="card-title">Compare choices</div><div class="card-sub">Understand trade-offs and next steps.</div></div>', unsafe_allow_html=True)
    with examples[2]:
        st.markdown('<div class="card"><div class="card-title">Prepare a plan</div><div class="card-sub">Turn a broad goal into weekly actions.</div></div>', unsafe_allow_html=True)
    for item in st.session_state.career_messages:
        with st.chat_message(item["role"]):
            st.markdown(item["content"])
    voice_question = ""
    audio = st.audio_input("Or record your question", sample_rate=16000)
    if audio and st.button("Transcribe and ask", disabled=not bool(API_KEY), key="transcribe_career_question"):
        if audio.size > 20 * 1024 * 1024:
            st.error("Recording is too large. Record a shorter question and try again.")
        else:
            try:
                with st.spinner("Transcribing your voice question…"):
                    transcript = client.audio.transcriptions.create(
                        file=(audio.name or "career-question.wav", audio.getvalue()),
                        model="whisper-large-v3-turbo",
                        response_format="json",
                        temperature=0.0,
                    )
                voice_question = (transcript.text or "").strip()
                if voice_question:
                    st.caption(f"Recognized: {voice_question}")
                else:
                    st.warning("I couldn't recognize speech. Try recording again or type your question.")
            except APIError as e:
                show_error(e)
            except Exception as e:
                st.error(f"Voice transcription failed: {e}")
    question = st.chat_input("Ask about career paths, role fit, or a decision…", disabled=not bool(API_KEY), max_chars=2500)
    question = question or voice_question
    if question:
        st.session_state.career_messages.append({"role": "user", "content": question})
        with st.chat_message("user"):
            st.markdown(question)
        system = f"""You are CareerSpring, a supportive career exploration assistant for students and job seekers. Reply in {language}. Hinglish means Hindi in Roman letters mixed naturally with English; Hindi uses Devanagari. Use the profile as context, not as a basis for stereotypes. Ask concise follow-up questions when needed. Offer 2-4 options with fit rationale, trade-offs, practical next steps, and ways to test fit. Do not guarantee jobs, salaries, admissions, or outcomes. Do not invent live openings, salary figures, demand statistics, or course availability; say when data needs current verification. Avoid claiming a formal personality or psychometric test. Keep answer readable and under 450 words. Profile:\n{profile}"""
        try:
            with st.chat_message("assistant"):
                with st.spinner("Thinking through your options…"):
                    answer = complete(system, question, st.session_state.career_messages[:-1])
                st.markdown(answer)
            st.session_state.career_messages.append({"role": "assistant", "content": answer})
        except APIError as e:
            st.session_state.career_messages.pop()
            show_error(e)
        except Exception as e:
            st.session_state.career_messages.pop()
            st.error(str(e))

with tab_resume:
    st.subheader("Get actionable resume feedback")
    st.caption("Upload PDF, DOCX, or TXT. The prototype checks content and structure; it does not run a certified ATS or guarantee recruiter outcomes.")
    job_desc = st.text_area("Target role or job description (optional)", height=120, placeholder="Paste the role requirements to tailor the review…")
    resume_file = st.file_uploader("Choose your resume", type=["pdf", "docx", "txt"], key="resume_file")
    if resume_file:
        if resume_file.size > 5 * 1024 * 1024:
            st.error("File is too large. Use a file under 5 MB.")
        else:
            st.success(f"Ready to review: {resume_file.name} ({resume_file.size // 1024} KB)")
            if st.button("Analyze my resume", type="primary", disabled=not bool(API_KEY)):
                try:
                    raw = resume_file.getvalue()
                    suffix = Path(resume_file.name).suffix.lower()
                    if suffix == ".pdf":
                        if PdfReader is None:
                            raise RuntimeError("PDF reader not installed. Check requirements.txt.")
                        import io
                        reader = PdfReader(io.BytesIO(raw))
                        text = "\n".join(page.extract_text() or "" for page in reader.pages)
                    elif suffix == ".docx":
                        if Document is None:
                            raise RuntimeError("DOCX reader not installed. Check requirements.txt.")
                        import io
                        doc = Document(io.BytesIO(raw))
                        text = "\n".join(p.text for p in doc.paragraphs)
                    else:
                        text = raw.decode("utf-8", errors="replace")
                    if len(text.strip()) < 80:
                        raise RuntimeError("Not enough selectable text found. If this is a scanned PDF, try a text-based PDF or paste the resume text into a TXT file.")
                    with st.spinner("Reviewing resume structure and impact…"):
                        feedback = complete(
                            f"You are a careful resume coach. Reply in {language}. Review only the supplied resume text and target role. Do not invent experience, credentials, skills, metrics, or ATS scores. Flag missing evidence as questions, distinguish high-priority fixes from optional edits, include 3 concrete before/after bullet rewrites using only known facts (use [add real metric] placeholders where needed), and finish with a short checklist. Mention that actual ATS systems vary. Keep within 700 words. Profile: {profile}",
                            f"TARGET ROLE / JOB DESCRIPTION:\n{job_desc or 'Not provided'}\n\nRESUME TEXT:\n{text[:18000]}",
                            max_tokens=1800,
                        )
                    st.markdown(feedback)
                except APIError as e:
                    show_error(e)
                except Exception as e:
                    st.error(str(e))

with tab_roadmap:
    st.subheader("Build a practical learning path")
    c1, c2 = st.columns(2)
    with c1:
        target_role = st.text_input("Target role", placeholder="e.g., Data Analyst")
        weekly_hours = st.slider("Hours available each week", 2, 20, 6)
    with c2:
        time_horizon = st.selectbox("Plan length", ["4 weeks", "8 weeks", "12 weeks"])
        budget = st.selectbox("Learning budget", ["Free resources preferred", "Low-cost options", "Any budget"])
    projects = st.text_area("What have you built or practiced so far?", height=100)
    if st.button("Create my roadmap", type="primary", disabled=not bool(API_KEY)):
        if not target_role:
            st.info("Add a target role first.")
        else:
            prompt = f"Create a realistic {time_horizon} beginner-friendly learning plan toward {target_role}. User profile: {profile}. Prior work: {projects or 'not provided'}. Time: {weekly_hours} hours/week. Budget: {budget}. Reply in {language}. Include a brief starting-level assumption, weekly milestones, one portfolio project that can be scoped to the time available, self-checks, and a final readiness checklist. Recommend only well-known learning platforms/topics; do not invent current course links, prices, certificates, or availability. Mark items to verify. Do not promise employment."
            try:
                with st.spinner("Building your plan…"):
                    st.markdown(complete("You are a practical learning-path coach. Keep the plan concrete, achievable, and honest about uncertainty.", prompt, max_tokens=1800))
            except APIError as e:
                show_error(e)
            except Exception as e:
                st.error(str(e))

with tab_interview:
    st.subheader("Practice a role-specific interview")
    role = st.text_input("Role to practice", placeholder="e.g., Python intern", key="interview_role")
    level = st.selectbox("Interview level", ["Internship / beginner", "Entry level", "Intermediate"])
    for item in st.session_state.interview_messages:
        with st.chat_message(item["role"]):
            st.markdown(item["content"])
    if not st.session_state.interview_messages:
        if st.button("Start mock interview", type="primary", disabled=not bool(API_KEY)):
            if not role:
                st.info("Enter a target role first.")
            else:
                try:
                    q = complete(
                        "You are an interviewer. Ask exactly one concise, realistic opening question for the specified role and level. Do not include an answer or evaluation yet.",
                        f"Role: {role}. Level: {level}. Language: {language}.", max_tokens=250)
                    st.session_state.interview_messages = [{"role": "assistant", "content": q}]
                    st.rerun()
                except APIError as e:
                    show_error(e)
                except Exception as e:
                    st.error(str(e))
    else:
        reply = st.chat_input("Type your interview answer…", disabled=not bool(API_KEY), max_chars=3000, key="interview_input")
        if reply:
            st.session_state.interview_messages.append({"role": "user", "content": reply})
            try:
                answer = complete(
                    f"You are a supportive mock interviewer for {role or 'the target role'} at {level}. Reply in {language}. Give brief, specific feedback on clarity, relevance, and evidence. Do not infer voice confidence or sentiment from typed text. Then ask exactly one next interview question. Do not claim a validated score. Keep under 250 words.",
                    reply,
                    st.session_state.interview_messages[:-1],
                    max_tokens=700,
                )
                st.session_state.interview_messages.append({"role": "assistant", "content": answer})
                st.rerun()
            except APIError as e:
                st.session_state.interview_messages.pop()
                show_error(e)
            except Exception as e:
                st.session_state.interview_messages.pop()
                st.error(str(e))
        if st.button("End interview and get summary"):
            if len(st.session_state.interview_messages) < 2:
                st.info("Answer at least one question first.")
            else:
                transcript = "\n".join(f"{m['role']}: {m['content']}" for m in st.session_state.interview_messages)
                try:
                    with st.spinner("Preparing feedback…"):
                        st.markdown(complete(
                            f"You are a mock interview coach. Reply in {language}. Review the transcript and give 3 strengths, 3 areas to practice, and a simple next practice drill. Evaluate only the written answers; do not infer voice, facial expression, confidence, or personality. Do not present a validated score or hiring prediction. Under 350 words.",
                            transcript, max_tokens=900))
                except APIError as e:
                    show_error(e)
                except Exception as e:
                    st.error(str(e))

st.divider()
st.caption("CareerSpring is an educational prototype. It does not connect to live job portals, verify current salaries, run trained matching models, or replace a qualified career professional. Resume text and voice recordings are sent to Groq only when you request AI analysis or transcription.")
